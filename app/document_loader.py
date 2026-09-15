"""Turns arbitrary uploaded files into plain text.

Runtime document sets need runtime-flexible loading: this module inspects
the file extension and dispatches to the right extractor. Adding a new file
type is a matter of adding one function + one dict entry, no pipeline
changes required.
"""
from __future__ import annotations

import csv
import io
import os
from typing import Callable, Dict


class UnsupportedFileTypeError(ValueError):
    pass


def _load_txt_or_md(raw: bytes) -> str:
    return raw.decode("utf-8", errors="replace")


def _load_pdf(raw: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(raw))
    pages = []
    for page in reader.pages:
        text = page.extract_text() or ""
        pages.append(text)
    return "\n\n".join(pages)


def _load_docx(raw: bytes) -> str:
    import docx

    doc = docx.Document(io.BytesIO(raw))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def _load_csv(raw: bytes) -> str:
    text = raw.decode("utf-8", errors="replace")
    reader = csv.reader(io.StringIO(text))
    rows = [", ".join(row) for row in reader]
    return "\n".join(rows)


_LOADERS: Dict[str, Callable[[bytes], str]] = {
    ".txt": _load_txt_or_md,
    ".md": _load_txt_or_md,
    ".markdown": _load_txt_or_md,
    ".pdf": _load_pdf,
    ".docx": _load_docx,
    ".csv": _load_csv,
}

SUPPORTED_EXTENSIONS = tuple(_LOADERS.keys())


def load_document(filename: str, raw: bytes) -> str:
    """Extract plain text from a document's raw bytes based on its extension.

    Raises UnsupportedFileTypeError for anything not in SUPPORTED_EXTENSIONS.
    """
    ext = os.path.splitext(filename)[1].lower()
    loader = _LOADERS.get(ext)
    if loader is None:
        raise UnsupportedFileTypeError(
            f"'{ext or filename}' is not supported. "
            f"Supported types: {', '.join(SUPPORTED_EXTENSIONS)}"
        )
    text = loader(raw)
    return text.strip()
